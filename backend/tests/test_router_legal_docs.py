"""
Integration tests for the LegalDocsRouter (/legal-docs): court list, text
extraction and PDF conversion.
"""

import asyncio
import io
from unittest.mock import AsyncMock, patch

import pytest

from tests.helpers import headers, make_user

pytestmark = pytest.mark.usefixtures("no_rate_limit")

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("Master Services Agreement", level=1)
    doc.add_paragraph("Either party may terminate on ninety (90) days written notice.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


class TestConvertToPdf:
    """The document viewer renders DOCX and text files through this endpoint;
    it depends on reportlab, which must be installed (no importorskip here)."""

    def test_requires_auth(self, client):
        resp = client.post(
            "/api/v1/legal-docs/convert-to-pdf",
            files={"file": ("notes.txt", b"Some notes.", "text/plain")},
        )
        assert resp.status_code == 401

    def test_docx_is_rendered_to_pdf(self, client, auth_headers):
        resp = client.post(
            "/api/v1/legal-docs/convert-to-pdf",
            files={"file": ("agreement.docx", _docx_bytes(), DOCX_MIME)},
            headers=auth_headers,
        )
        assert resp.status_code == 200
        assert resp.headers["content-type"] == "application/pdf"
        assert resp.content[:4] == b"%PDF"

    def test_text_is_rendered_to_pdf(self, client, auth_headers):
        resp = client.post(
            "/api/v1/legal-docs/convert-to-pdf",
            files={"file": ("notes.txt", b"First line.\n\nSecond paragraph.", "text/plain")},
            headers=auth_headers,
        )
        assert resp.status_code == 200
        assert resp.content[:4] == b"%PDF"

    def test_pdf_passes_through_unchanged(self, client, auth_headers):
        pdf = b"%PDF-1.4\n1 0 obj\n<<>>\nendobj\ntrailer\n<<>>\n%%EOF\n"
        resp = client.post(
            "/api/v1/legal-docs/convert-to-pdf",
            files={"file": ("filing.pdf", pdf, "application/pdf")},
            headers=auth_headers,
        )
        assert resp.status_code == 200
        assert resp.content == pdf


class TestExtractText:
    def test_requires_auth(self, client):
        resp = client.post(
            "/api/v1/legal-docs/extract-text",
            files={"file": ("notes.txt", b"Some notes.", "text/plain")},
        )
        assert resp.status_code == 401

    def test_extracts_text_off_the_event_loop(self, client, auth_headers):
        """Extraction is OCR-capable and slow; it must run in a worker thread."""
        from app.services.text_extraction import text_extraction_service

        real_extract = text_extraction_service.extract
        on_event_loop = []

        def _spy(**kwargs):
            try:
                asyncio.get_running_loop()
                on_event_loop.append(True)
            except RuntimeError:  # no loop in this thread -> a worker thread
                on_event_loop.append(False)
            return real_extract(**kwargs)

        with patch.object(text_extraction_service, "extract", side_effect=_spy):
            resp = client.post(
                "/api/v1/legal-docs/extract-text",
                files={"file": ("notes.txt", b"First line.\nSecond line.", "text/plain")},
                headers=auth_headers,
            )
        assert resp.status_code == 200
        data = resp.json()
        assert data["success"] is True
        assert "First line." in data["text"]
        assert data["word_count"] == 4
        assert on_event_loop == [False]

    def test_file_with_no_extractable_text_reports_the_reason(self, client, auth_headers):
        from app.services.text_extraction import ExtractionResult, text_extraction_service

        reason = "No text could be extracted. This looks like a scanned PDF."
        with patch.object(
            text_extraction_service, "extract", return_value=ExtractionResult(error=reason)
        ):
            resp = client.post(
                "/api/v1/legal-docs/extract-text",
                files={"file": ("notes.txt", b"Some notes.", "text/plain")},
                headers=auth_headers,
            )
        assert resp.status_code == 422
        assert resp.json()["detail"] == reason

    def test_disallowed_file_type_is_rejected(self, client, auth_headers):
        resp = client.post(
            "/api/v1/legal-docs/extract-text",
            files={"file": ("run.exe", b"MZ\x90\x00binary", "application/x-msdownload")},
            headers=auth_headers,
        )
        assert resp.status_code in (400, 415)


class TestPermissionGate:
    """Extraction/conversion need a document, contract or authority-map permission."""

    @pytest.mark.parametrize("endpoint", ["extract-text", "convert-to-pdf"])
    def test_user_without_any_workspace_permission_is_forbidden(self, client, endpoint):
        user = make_user("ld-user")
        with patch("app.services.permissions.effective_permissions", AsyncMock(return_value=set())):
            resp = client.post(
                f"/api/v1/legal-docs/{endpoint}",
                files={"file": ("notes.txt", b"Some notes.", "text/plain")},
                headers=headers(user),
            )
        assert resp.status_code == 403

    def test_viewer_can_use_the_document_viewer_conversion(self, client, non_admin_headers):
        """Viewers hold documents.view, which the in-app document viewer relies on."""
        pdf = b"%PDF-1.4\n%%EOF\n"
        resp = client.post(
            "/api/v1/legal-docs/convert-to-pdf",
            files={"file": ("filing.pdf", pdf, "application/pdf")},
            headers=non_admin_headers,
        )
        assert resp.status_code == 200

    def test_courts_requires_auth_only(self, client, non_admin_headers):
        assert client.get("/api/v1/legal-docs/courts").status_code == 401
        resp = client.get("/api/v1/legal-docs/courts", headers=non_admin_headers)
        assert resp.status_code == 200
        assert resp.json()["count"] == len(resp.json()["courts"]) > 0
