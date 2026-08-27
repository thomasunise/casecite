"""
Tests for conversation export rendering.

Source: backend/app/services/conversation_export.py
"""

import io
import os

import pytest

os.environ["SECRET_KEY"] = "test-secret-key-for-testing-only-32chars!"
os.environ["ENCRYPTION_SALT"] = "test-salt-16chars!"
os.environ["DEBUG"] = "true"

from app.services.conversation_export import (  # noqa: E402
    parse_blocks,
    render_docx,
    render_markdown,
    render_pdf,
    safe_filename,
    strip_inline,
)

MESSAGES = [
    {"role": "user", "text": "What does the lease say about **early termination**?"},
    {
        "role": "assistant",
        "text": (
            "## Early termination\n\n"
            "The tenant may terminate with **90 days** notice.\n\n"
            "- Notice must be in writing\n"
            "- A fee of `two months rent` applies\n\n"
            "1. Serve notice\n"
            "2. Pay the fee\n"
        ),
        "citations": [
            {"label": "Lease.pdf", "quote": "ninety (90) days written notice"},
            {
                "label": "Smith v. Jones, 12 F.3d 34",
                "url": "https://www.courtlistener.com/opinion/1/",
            },
        ],
    },
]


class TestParsing:
    def test_strip_inline_markers(self):
        assert (
            strip_inline("**bold** and _it_ and `code` and [x](http://y)")
            == "bold and it and code and x"
        )

    def test_blocks(self):
        kinds = [k for k, _ in parse_blocks(MESSAGES[1]["text"])]
        assert kinds == ["heading", "para", "bullet", "bullet", "number", "number"]
        assert parse_blocks("")[0:0] == []


class TestMarkdown:
    def test_structure_and_sources(self):
        md = render_markdown("Lease questions", MESSAGES)
        assert md.startswith("# Lease questions")
        assert "## You" in md and "## Assistant" in md
        assert "**Sources**" in md
        assert "1. Lease.pdf" in md
        assert "> ninety (90) days written notice" in md
        assert "courtlistener.com" in md

    def test_empty_title_defaults(self):
        assert render_markdown("", []).startswith("# Conversation")


class TestDocx:
    def test_docx_contains_text_without_markdown_markers(self):
        from docx import Document

        doc = Document(io.BytesIO(render_docx("Lease questions", MESSAGES)))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Lease questions" in text
        assert "The tenant may terminate with 90 days notice." in text
        assert "**" not in text and "`" not in text
        assert "Notice must be in writing" in text
        assert "ninety (90) days written notice" in text
        styles = [p.style.name for p in doc.paragraphs]
        assert "List Bullet" in styles and "List Number" in styles


class TestPdf:
    def test_pdf_bytes(self):
        pytest.importorskip("reportlab")
        blob = render_pdf("Lease questions", MESSAGES)
        assert blob[:4] == b"%PDF"
        assert len(blob) > 500


class TestFilename:
    def test_safe_filename(self):
        assert safe_filename("Smith v. Jones — 2026", "docx") == "Smith-v-Jones-2026.docx"
        assert safe_filename("", "pdf") == "conversation.pdf"
        assert safe_filename("é•ü", "md") == "conversation.md"


class TestReviewNotice:
    """Every export format ends with the AI review notice."""

    def test_markdown(self):
        from app.services.ai_notice import AI_REVIEW_NOTICE

        md = render_markdown("Lease questions", MESSAGES)
        assert md.rstrip().endswith(f"_{AI_REVIEW_NOTICE}_")
        assert md.rstrip().splitlines()[-3] == "---"

    def test_docx(self):
        from app.services.ai_notice import AI_REVIEW_NOTICE
        from docx import Document

        doc = Document(io.BytesIO(render_docx("Lease questions", MESSAGES)))
        last = doc.paragraphs[-1]
        assert last.text == AI_REVIEW_NOTICE
        assert all(run.italic for run in last.runs)

    def test_pdf(self):
        import base64
        import re
        import zlib

        pytest.importorskip("reportlab")
        blob = render_pdf("Lease questions", MESSAGES)
        # reportlab content streams are ASCII85 + Flate encoded.
        text = b""
        for chunk in re.findall(rb"stream(.*?)endstream", blob, re.S):
            chunk = chunk.strip()
            if chunk.endswith(b"~>"):
                chunk = chunk[:-2]
            try:
                text += zlib.decompress(base64.a85decode(chunk))
            except (ValueError, zlib.error):
                text += chunk
        assert b"AI-generated draft." in text
        assert b"before relying on it or sending it to anyone." in text
