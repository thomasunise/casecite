"""Document drafting: strict-JSON output with a plain-text fallback, docx render."""

import io
import json
from unittest.mock import AsyncMock, MagicMock

import pytest
from app.services.contract_analysis.drafting import draft_document, render_draft_docx


def _client_returning(content: str) -> MagicMock:
    client = MagicMock()
    response = MagicMock()
    response.choices = [MagicMock()]
    response.choices[0].message.content = content
    client.chat.completions.create = AsyncMock(return_value=response)
    return client


class TestDraftDocument:
    @pytest.mark.asyncio
    async def test_json_payload_yields_title_and_text(self):
        payload = {
            "title": "First Amendment to Master Services Agreement",
            "text": "FIRST AMENDMENT\n\nThe Term is extended to [DATE].",
        }
        client = _client_returning(json.dumps(payload))
        result = await draft_document(client, "Draft an amendment extending the term")
        assert result["title"] == payload["title"]
        assert result["text"] == payload["text"]

    @pytest.mark.asyncio
    async def test_non_json_response_falls_back_to_draft_title(self):
        raw = "DEMAND LETTER\n\nDear [RECIPIENT],\n\nPlease remit payment of [AMOUNT]."
        client = _client_returning(raw)
        result = await draft_document(client, "Draft a demand letter")
        assert result["title"] == "Draft"
        assert result["text"] == raw

    @pytest.mark.asyncio
    async def test_reference_text_is_capped_by_the_model_window(self, monkeypatch):
        from app.services.contract_analysis import drafting

        # A model with a tiny window: the cap falls back to the 40k floor.
        monkeypatch.setattr(drafting, "context_window_tokens", lambda: 8_000)
        client = _client_returning(json.dumps({"title": "Letter", "text": "Body"}))
        await draft_document(client, "Draft a response letter", reference_text="X" * 50_000)
        kwargs = client.chat.completions.create.call_args.kwargs
        user_msg = kwargs["messages"][-1]["content"]
        assert "REFERENCE DOCUMENT" in user_msg
        assert user_msg.count("X") == 40_000

    @pytest.mark.asyncio
    async def test_reference_text_is_not_truncated_on_a_large_window(self, monkeypatch):
        from app.services.contract_analysis import drafting

        monkeypatch.setattr(drafting, "context_window_tokens", lambda: 400_000)
        client = _client_returning(json.dumps({"title": "Letter", "text": "Body"}))
        await draft_document(client, "Draft a response letter", reference_text="X" * 120_000)
        user_msg = client.chat.completions.create.call_args.kwargs["messages"][-1]["content"]
        assert user_msg.count("X") == 120_000

    @pytest.mark.asyncio
    async def test_none_client_raises_value_error(self):
        with pytest.raises(ValueError):
            await draft_document(None, "Draft anything")

    @pytest.mark.asyncio
    async def test_llm_failure_raises_runtime_error(self):
        client = MagicMock()
        client.chat.completions.create = AsyncMock(side_effect=ConnectionError("down"))
        with pytest.raises(RuntimeError):
            await draft_document(client, "Draft a letter")

    @pytest.mark.asyncio
    async def test_empty_response_raises_runtime_error(self):
        client = _client_returning("")
        with pytest.raises(RuntimeError):
            await draft_document(client, "Draft a letter")


class TestRenderDraftDocx:
    def test_returns_pk_zip_bytes(self):
        data = render_draft_docx("My Draft", "Line one.\n\nLine two.")
        assert isinstance(data, bytes)
        assert data[:2] == b"PK"

    def test_round_trips_through_document(self):
        from docx import Document

        title = "Settlement Agreement"
        text = "First paragraph.\n\nSecond paragraph."
        data = render_draft_docx(title, text)
        doc = Document(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs]
        assert paragraphs[0] == title
        assert doc.paragraphs[0].style.name.startswith("Heading 1")
        assert "First paragraph." in paragraphs
        assert "Second paragraph." in paragraphs
        # The blank line between paragraphs survives as an empty paragraph.
        assert "" in paragraphs[1:]


class TestDraftCaseLawGuard:
    """A draft is model output: it may not introduce authority of its own."""

    @pytest.mark.asyncio
    async def test_invented_citation_is_removed_and_reported(self):
        payload = {
            "title": "Motion",
            "text": "ARGUMENT\n\nSee Madeup v. Imaginary, 999 F.3d 1. Relief is warranted.",
        }
        result = await draft_document(_client_returning(json.dumps(payload)), "Draft a motion")
        assert "Madeup" not in result["text"]
        assert "[unverified case-law reference removed]" in result["text"]
        assert "Relief is warranted." in result["text"]
        assert len(result["case_law_removed"]) == 1

    @pytest.mark.asyncio
    async def test_case_named_in_the_instructions_is_kept(self):
        payload = {"title": "Motion", "text": "As held in Smith v. Jones, notice is required."}
        result = await draft_document(
            _client_returning(json.dumps(payload)),
            "Draft a motion relying on Smith v. Jones",
        )
        assert result["text"] == payload["text"]
        assert result["case_law_removed"] == []

    @pytest.mark.asyncio
    async def test_case_in_the_reference_document_or_current_draft_is_kept(self):
        payload = {"title": "Brief", "text": "Under Roe v. Wade, 410 U.S. 113, the rule applies."}
        client = _client_returning(json.dumps(payload))
        from_reference = await draft_document(
            client, "Draft a brief", reference_text="Our memo discusses Roe v. Wade, 410 U.S. 113."
        )
        assert from_reference["case_law_removed"] == []
        from_draft = await draft_document(
            client, "Tighten the prose", revision_of="Under Roe v. Wade, 410 U.S. 113, it applies."
        )
        assert from_draft["case_law_removed"] == []

    @pytest.mark.asyncio
    async def test_contract_draft_without_case_references_is_untouched(self):
        payload = {"title": "NDA", "text": "1. CONFIDENTIALITY\n\nEach party shall keep it secret."}
        result = await draft_document(_client_returning(json.dumps(payload)), "Draft an NDA")
        assert result["text"] == payload["text"]
        assert result["case_law_removed"] == []
