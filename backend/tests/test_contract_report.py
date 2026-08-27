"""
Unit tests for the contract-analysis report renderer
(app/services/contract_analysis/report.py). Pure rendering — no DB, no HTTP.
"""

import pytest
from app.services.ai_notice import AI_REVIEW_NOTICE
from app.services.contract_analysis.report import render_docx, render_markdown


@pytest.fixture
def data():
    return {
        "analysis_id": "abc12345-dead-beef-0000-000000000000",
        "contract_type": "msa",
        "created_at": "2026-08-18T14:30:00",
        "representing": "Customer",
        "posture": "strict",
        "executive_summary": "A one-sided master services agreement favoring the vendor.",
        "key_terms": {
            "governing_law": {
                "value": "New York",
                "quote": "This Agreement shall be governed by the laws of the State of New York.",
                "span_start": 100,
                "span_end": 170,
                "verified": True,
            },
            "liability_cap": {
                "value": "12 months of fees",
                "quote": "not found verbatim",
                "span_start": None,
                "span_end": None,
                "verified": False,
            },
        },
        "issues": [
            {
                "ref": "dev-0",
                "title": "Uncapped indemnification",
                "status": "off_market",
                "severity": "critical",
                "why": "Customer bears unlimited indemnity exposure.",
                "suggested_language": "Indemnity is capped at fees paid in the prior 12 months.",
                "span_start": 500,
                "span_end": 600,
                "matched_text": "Customer shall indemnify Vendor against all claims.",
                "grounding": "span",
            },
            {
                "ref": "miss-0",
                "title": "Missing clause: Limitation of Liability",
                "status": "missing",
                "severity": "major",
                "why": "No liability cap protects the Customer.",
                "suggested_language": None,
                "span_start": None,
                "span_end": None,
                "grounding": "absence",
            },
            {
                "ref": "gen-0",
                "title": "Vague acceptance criteria",
                "status": "unusual",
                "severity": "minor",
                "why": "Deliverable acceptance is undefined.",
                "suggested_language": None,
                "span_start": None,
                "span_end": None,
                "grounding": "unverified",
            },
        ],
        "parties": [
            {"canonical_name": "Acme Corp", "role": "vendor", "aliases": ["Acme"]},
            {"canonical_name": "Widget LLC", "role": "customer", "aliases": []},
        ],
        "obligations": [
            {
                "subject_party": "Widget LLC",
                "modal": "shall",
                "action": "pay all invoices",
                "object_text": "within thirty (30) days of receipt",
                "deadlines": [{"description": "within 30 days of invoice receipt"}],
            },
            {
                "subject_party": "Acme Corp",
                "modal": "shall",
                "action": "maintain insurance",
                "object_text": "throughout the Term",
                "deadlines": [],
            },
        ],
        "deadlines": [
            {
                "description": "Payment due within 30 days of invoice",
                "matched_text": "within thirty (30) days",
                "resolved_date": "2026-09-17T00:00:00",
            },
            {
                "description": "Renewal notice 60 days before term end",
                "matched_text": None,
                "resolved_date": None,
            },
        ],
    }


class TestRenderMarkdown:
    def test_header_and_metadata(self, data):
        md = render_markdown(data)
        assert md.startswith("# Contract Analysis Report")
        assert "msa" in md
        assert "2026-08-18" in md
        assert "Customer" in md
        assert "strict" in md

    def test_executive_summary(self, data):
        md = render_markdown(data)
        assert "## Executive Summary" in md
        assert "one-sided master services agreement" in md

    def test_key_terms_verified_quote_and_unverified_tag(self, data):
        md = render_markdown(data)
        assert "## Key Terms" in md
        assert "Governing Law" in md
        assert "New York" in md
        assert "governed by the laws of the State of New York" in md
        # Unverified value is tagged and its unlocatable quote is not presented as sourced.
        assert "12 months of fees (unverified)" in md
        assert '"not found verbatim"' not in md

    def test_issues_grouped_by_severity_in_order(self, data):
        md = render_markdown(data)
        assert "## Issues" in md
        crit = md.index("### Critical")
        major = md.index("### Major")
        minor = md.index("### Minor")
        assert crit < major < minor
        assert "Uncapped indemnification" in md
        assert "Customer bears unlimited indemnity exposure." in md
        assert "Indemnity is capped at fees paid in the prior 12 months." in md
        assert "Customer shall indemnify Vendor against all claims." in md

    def test_absence_and_unverified_issue_tags(self, data):
        md = render_markdown(data)
        assert "Clause not found in document." in md
        assert "Vague acceptance criteria (unverified — model conclusion)" in md

    def test_parties(self, data):
        md = render_markdown(data)
        assert "## Parties" in md
        assert "Acme Corp (vendor)" in md
        assert "Widget LLC (customer)" in md

    def test_obligations_with_linked_deadlines(self, data):
        md = render_markdown(data)
        assert "## Obligations" in md
        assert "Widget LLC — pay all invoices — within thirty (30) days of receipt" in md
        assert "Deadline: within 30 days of invoice receipt" in md
        # Obligation with no linked deadlines renders without a deadline sub-item.
        assert "Acme Corp — maintain insurance — throughout the Term" in md

    def test_deadlines_section(self, data):
        md = render_markdown(data)
        assert "## Deadlines" in md
        assert "Payment due within 30 days of invoice — resolved date: 2026-09-17" in md
        assert "Renewal notice 60 days before term end" in md

    def test_no_risk_score_in_report(self, data):
        md = render_markdown(data)
        assert "Risk Score" not in md
        assert "risk score" not in md.lower()

    def test_minimal_data_does_not_crash(self):
        md = render_markdown({"analysis_id": "x"})
        assert "# Contract Analysis Report" in md


class TestRenderDocx:
    def test_returns_zip_bytes(self, data):
        blob = render_docx(data)
        assert isinstance(blob, bytes)
        assert len(blob) > 0
        assert blob[:2] == b"PK"

    def test_document_contains_report_text(self, data):
        import io

        from docx import Document

        doc = Document(io.BytesIO(render_docx(data)))
        text = "\n".join(p.text for p in doc.paragraphs)
        table_text = "\n".join(
            cell.text for t in doc.tables for row in t.rows for cell in row.cells
        )
        assert "Contract Analysis Report" in text
        assert "Uncapped indemnification" in text
        assert "Clause not found in document." in text
        assert "(unverified — model conclusion)" in text
        assert "Governing Law" in table_text
        assert "12 months of fees (unverified)" in table_text

    def test_minimal_data_does_not_crash(self):
        blob = render_docx({"analysis_id": "x"})
        assert blob[:2] == b"PK"


class TestReviewNotice:
    """Every exported report ends with the AI review notice."""

    def test_markdown_ends_with_notice_after_rule(self, data):
        md = render_markdown(data)
        assert md.rstrip().endswith(f"_{AI_REVIEW_NOTICE}_")
        tail = md.rstrip().splitlines()[-3:]
        assert tail[0] == "---"

    def test_docx_last_paragraph_is_italic_notice(self, data):
        import io

        from docx import Document

        doc = Document(io.BytesIO(render_docx(data)))
        last = doc.paragraphs[-1]
        assert last.text == AI_REVIEW_NOTICE
        assert all(run.italic for run in last.runs)

    def test_minimal_report_still_carries_notice(self):
        assert AI_REVIEW_NOTICE in render_markdown({"analysis_id": "x"})
