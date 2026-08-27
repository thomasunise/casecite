"""Document drafting from natural-language instructions.

One call to the full chat model (drafting is generative work, not extraction —
the utility model is deliberately not used here) with the same no-preamble
drafting rules as the RAG drafting path (``app/services/rag/prompts.py``). The
currently selected contract can be passed as reference/style material. Output
is strict JSON {"title", "text"}; a model that ignores the JSON instruction
still yields a usable draft via the plain-text fallback.
"""

from __future__ import annotations

import io
import json
import logging

from app.services.contract_analysis.long_drafting import context_window_tokens
from app.services.llm_clients import chat_model, openai_chat
from app.services.rag.prompt_safety import UNTRUSTED_CONTENT_RULE, untrusted_block

logger = logging.getLogger(__name__)

# Never send less than this much reference/draft text, whatever the model claims.
REFERENCE_CHARS_FLOOR = 40_000
# Output tokens kept free for the answer (a revision returns the whole document).
_OUTPUT_RESERVE_TOKENS = 40_000


def reference_cap_chars() -> int:
    """Characters of reference or current-draft text one call may carry.

    Bounded by the configured model's context window, not a fixed constant —
    a fixed 40k-character cut silently dropped everything past page ten of a
    long reference contract (or of the draft being revised).
    """
    return max(REFERENCE_CHARS_FLOOR, (context_window_tokens() - _OUTPUT_RESERVE_TOKENS) * 4)


# Mirrors the is_drafting system prompt in app/services/rag/prompts.py, plus
# the strict-JSON envelope this endpoint needs.
_DRAFTING_SYSTEM_PROMPT = (
    """You are an expert legal document drafter. Your role is to write complete, properly formatted legal documents.

IMPORTANT RULES:
- Write the document body only — no preamble, commentary, or explanation before or after it (the product adds its own review notice)
- Use [BRACKETED PLACEHOLDERS] for facts you don't know
- Write professionally formatted legal documents following standard conventions
- Start directly with the document content (caption, title, etc.)

Return STRICT JSON: {"title": "<short document title>", "text": "<the full document text>"}.

"""
    + UNTRUSTED_CONTENT_RULE
)


async def draft_document(
    client,
    instructions: str,
    reference_text: str | None = None,
    revision_of: str | None = None,
) -> dict:
    """Draft a contract/letter from instructions, optionally styled on a reference.

    ``revision_of`` switches to revision mode: the instructions are applied to
    that existing draft and the FULL revised document comes back — this is how
    the drafting workspace iterates ("make the term 3 years", "add arbitration").

    Returns {"title": str, "text": str}.

    Raises ValueError when ``client`` is None; any other failure raises
    RuntimeError with a clean message (the caller maps it to HTTP).
    """
    if client is None:
        raise ValueError("No LLM client available for document drafting")

    cap = reference_cap_chars()
    if revision_of and revision_of.strip():
        user_content = (
            "CURRENT DRAFT:\n"
            + untrusted_block("Current draft", revision_of[:cap])
            + f"\n\nREVISION INSTRUCTIONS:\n{instructions.strip()}\n\n"
            "Apply the revision instructions to the current draft and return the "
            "FULL revised document — every unchanged section repeated verbatim, "
            "never a summary of changes."
        )
    else:
        user_content = f"INSTRUCTIONS:\n{instructions.strip()}"
        if reference_text and reference_text.strip():
            user_content += (
                "\n\nREFERENCE DOCUMENT (use for context, defined terms, party names, "
                "and style — the draft should be consistent with it):\n"
                + untrusted_block("Reference document", reference_text[:cap])
            )

    try:
        resp = await openai_chat(
            client,
            model=chat_model(),
            messages=[
                {"role": "system", "content": _DRAFTING_SYSTEM_PROMPT},
                {"role": "user", "content": user_content},
            ],
            temperature=0.3,
            response_format={"type": "json_object"},
        )
        content = resp.choices[0].message.content or ""
    except Exception as e:  # surfaced to the caller as one clean error
        logger.warning("Document drafting failed: %s", e)
        raise RuntimeError(f"Document drafting failed: {e}") from e

    try:
        data = json.loads(content)
        title = str(data.get("title") or "").strip() or "Draft"
        text = str(data.get("text") or "").strip()
        if not text:
            raise ValueError("empty text field")
    except (json.JSONDecodeError, AttributeError, ValueError, TypeError):
        # Model ignored the JSON envelope — the whole response IS the draft.
        title = "Draft"
        text = content.strip()

    if not text:
        raise RuntimeError("Document drafting failed: the model returned an empty draft")
    return {"title": title, "text": text}


def render_draft_docx(title: str, text: str) -> bytes:
    """Render a draft as a .docx: Heading 1 title, then the text as paragraphs.

    Blank lines in the text are preserved as empty paragraphs so the document
    keeps its visual spacing.
    """
    from docx import Document

    from app.services.ai_notice import add_docx_notice

    doc = Document()
    doc.add_heading(title or "Draft", level=1)
    for line in (text or "").split("\n"):
        doc.add_paragraph(line.rstrip("\r"))
    add_docx_notice(doc)  # every generated artifact leaves with the review notice
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
