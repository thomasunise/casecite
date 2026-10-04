"""
Unified text extraction service.

Formats: PDF (tables + OCR fallback), DOCX, legacy DOC (Word 97-2003 binary),
XLSX/XLS, PPTX, ODT, CSV/TSV, HTML, JSON, EML/MSG email, RTF, TXT/Markdown.
Tables are preserved as markdown; every extractor degrades gracefully when its
optional library is missing.

Uploads are untrusted. ``extract`` never raises on a bad file: it returns an
``ExtractionResult`` whose ``error`` carries a reason a user can act on, and
whose ``warnings`` record anything that was skipped (page caps, OCR caps).
"""

import asyncio
import csv
import io
import logging
import zipfile
from dataclasses import dataclass, field
from datetime import UTC
from html.parser import HTMLParser

from app.config import settings

logger = logging.getLogger(__name__)

# Limits on untrusted input. A zip-based office file is tiny on the wire and
# arbitrarily large once inflated, and a PDF can declare any number of pages;
# both are parsed in the API process.
_MAX_UNCOMPRESSED_BYTES = getattr(settings, "extraction_max_uncompressed_bytes", 500 * 1024 * 1024)
_MAX_ZIP_MEMBERS = getattr(settings, "extraction_max_zip_members", 10_000)
_MAX_PDF_PAGES = getattr(settings, "extraction_max_pdf_pages", 5_000)
_MAX_OCR_PAGES = getattr(settings, "extraction_max_ocr_pages", 300)
# Longest edge, in pixels, of a page rendered for OCR (300 DPI Letter is 3300).
_MAX_OCR_PIXELS = 6_000


class ExtractionError(ValueError):
    """A file could not be extracted; the message is safe to show the user."""


# Optional: pdfplumber for table-aware PDF extraction
try:
    import pdfplumber

    PDFPLUMBER_AVAILABLE = True
except ImportError:
    PDFPLUMBER_AVAILABLE = False

# Fallback: pypdf for basic PDF text extraction
try:
    from pypdf import PdfReader

    PYPDF_AVAILABLE = True
except ImportError:
    PYPDF_AVAILABLE = False

# Optional: OCR support for scanned PDFs
try:
    import fitz  # PyMuPDF
    import pytesseract
    from PIL import Image

    OCR_AVAILABLE = True
except ImportError:
    OCR_AVAILABLE = False
    logger.info(
        "OCR disabled: PyMuPDF and/or pytesseract not installed — scanned-PDF "
        "pages will be indexed without OCR. Install PyMuPDF to enable it "
        "(`pip install PyMuPDF`); it is AGPL-licensed and therefore not "
        "bundled with this ELv2-licensed product."
    )

# DOCX support
try:
    from docx import Document as DocxDocument

    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

# XLSX support
try:
    import openpyxl

    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

# Legacy XLS support
try:
    import xlrd

    XLRD_AVAILABLE = True
except ImportError:
    XLRD_AVAILABLE = False

# PPTX support
try:
    from pptx import Presentation

    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

# Outlook .msg support (OLE container parsed directly — olefile is a single
# pure-python wheel, unlike the extract-msg dependency tree)
try:
    import olefile

    MSG_AVAILABLE = True
except ImportError:
    MSG_AVAILABLE = False

# RTF support
try:
    from striprtf.striprtf import rtf_to_text

    RTF_AVAILABLE = True
except ImportError:
    RTF_AVAILABLE = False


@dataclass
class ExtractionResult:
    text: str = ""
    tables: list[str] = field(default_factory=list)
    page_count: int = 0
    ocr_used: bool = False
    char_count: int = 0
    word_count: int = 0
    # Why no (or partial) text came back, in words a user can act on.
    error: str = ""
    warnings: list[str] = field(default_factory=list)
    truncated: bool = False


class TextExtractionService:
    """Unified text extraction with table preservation and OCR fallback."""

    # Minimum chars on a page before OCR is attempted
    OCR_THRESHOLD = 50

    def extract(
        self,
        file_path: str | None = None,
        file_content: bytes | None = None,
        content_type: str = "",
        filename: str = "",
    ) -> ExtractionResult:
        """
        Extract text from a file.

        Provide either file_path (reads from disk) or file_content (raw bytes).
        content_type or filename is used to determine format.
        """
        if file_content is None and file_path is not None:
            with open(file_path, "rb") as f:
                file_content = f.read()
        elif file_content is None:
            return ExtractionResult(error="No file content was provided.")

        # Determine format
        fmt = self._detect_format(file_content, content_type, filename)

        extractors = {
            "pdf": self._extract_pdf,
            "docx": self._extract_docx,
            "doc": self._extract_doc,
            "xlsx": self._extract_xlsx,
            "xls": self._extract_xls,
            "pptx": self._extract_pptx,
            "odt": self._extract_odt,
            "csv": lambda c: self._extract_delimited(c, ","),
            "tsv": lambda c: self._extract_delimited(c, "\t"),
            "html": self._extract_html,
            "eml": self._extract_eml,
            "msg": self._extract_msg,
            "rtf": self._extract_rtf,
        }
        # Unknown formats (and txt/md/json) fall through to plain text.
        extractor = extractors.get(fmt, self._extract_plain)
        try:
            result = extractor(file_content)
        except ExtractionError as e:
            result = ExtractionResult(error=str(e))
        except Exception as e:  # parser libraries raise their own exception trees
            # pdfminer, pypdf, python-docx, openpyxl and olefile each raise
            # library-specific errors on malformed input; one bad upload must
            # produce a failed document, not a 500.
            logger.error(f"{fmt} extraction failed: {type(e).__name__}: {e}")
            result = ExtractionResult(
                error=(
                    f"The file could not be read as {fmt.upper()}. It may be corrupt, "
                    "password-protected, or saved in a different format than its name suggests."
                )
            )

        if not result.text.strip() and not result.error:
            result.error = self._empty_reason(fmt)
        result.char_count = len(result.text)
        result.word_count = len(result.text.split()) if result.text.strip() else 0
        return result

    async def extract_async(
        self,
        file_path: str | None = None,
        file_content: bytes | None = None,
        content_type: str = "",
        filename: str = "",
    ) -> ExtractionResult:
        """``extract`` off the event loop. Parsing and OCR are CPU-bound and
        synchronous; async callers must use this so one large document does not
        stall every other request."""
        return await asyncio.to_thread(
            self.extract,
            file_path=file_path,
            file_content=file_content,
            content_type=content_type,
            filename=filename,
        )

    @staticmethod
    def _empty_reason(fmt: str) -> str:
        """Explain an extraction that produced no text."""
        if fmt == "pdf":
            if not OCR_AVAILABLE:
                return (
                    "No text could be extracted. This looks like a scanned (image-only) "
                    "PDF, and OCR is not enabled on this server. Upload a text-searchable "
                    "PDF, or ask your administrator to enable OCR."
                )
            return (
                "No text could be extracted, even with OCR. The pages may be blank or unreadable."
            )
        return "The file contains no extractable text."

    @staticmethod
    def _check_zip_container(content: bytes) -> None:
        """Reject zip-based office files that would inflate past the limits.

        Reads only the central directory, so the check itself is cheap. The
        declared sizes are what the parsers allocate from.
        """
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as zf:
                infos = zf.infolist()
        except (zipfile.BadZipFile, OSError, RuntimeError) as e:
            raise ExtractionError(
                "The file is not a valid Office document (it could not be opened as a "
                "zip container). It may be corrupt or saved in an older format."
            ) from e
        if len(infos) > _MAX_ZIP_MEMBERS:
            raise ExtractionError(
                f"The file contains too many internal parts ({len(infos):,}) to process safely."
            )
        total = sum(i.file_size for i in infos)
        if total > _MAX_UNCOMPRESSED_BYTES:
            raise ExtractionError(
                "The file expands to "
                f"{total / (1024 * 1024):,.0f} MB when decompressed, over the "
                f"{_MAX_UNCOMPRESSED_BYTES / (1024 * 1024):,.0f} MB limit."
            )

    def _detect_format(self, content: bytes, content_type: str, filename: str) -> str:
        """Detect file format from content type, filename, or magic bytes."""
        fn = filename.lower()

        by_ext = {
            ".pdf": "pdf",
            ".docx": "docx",
            ".doc": "doc",
            ".xlsx": "xlsx",
            ".xls": "xls",
            ".pptx": "pptx",
            ".odt": "odt",
            ".csv": "csv",
            ".tsv": "tsv",
            ".html": "html",
            ".htm": "html",
            ".eml": "eml",
            ".msg": "msg",
            ".rtf": "rtf",
        }
        by_mime = {
            "application/pdf": "pdf",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
            "application/msword": "doc",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
            "application/vnd.ms-excel": "xls",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
            "application/vnd.oasis.opendocument.text": "odt",
            "text/csv": "csv",
            "text/tab-separated-values": "tsv",
            "text/html": "html",
            "message/rfc822": "eml",
            "application/vnd.ms-outlook": "msg",
            "text/rtf": "rtf",
            "application/rtf": "rtf",
        }
        for ext, fmt in by_ext.items():
            if fn.endswith(ext):
                return fmt
        if content_type in by_mime:
            return by_mime[content_type]
        if content_type == "text/plain":
            return "txt"

        # Magic byte detection
        if content[:4] == b"%PDF":
            return "pdf"
        if content[:4] == b"PK\x03\x04":
            return self._classify_zip(content)
        if content[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
            # OLE compound file with no telling extension — legacy Word is the
            # common case; the doc extractor reports anything else clearly.
            return "doc"

        return "txt"

    def _classify_zip(self, content: bytes) -> str:
        """Distinguish the zip-container formats (docx/xlsx/pptx/odt) by their
        internal directory layout when the filename didn't tell us."""
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as zf:
                names = zf.namelist()
        except (zipfile.BadZipFile, OSError, RuntimeError):
            return "docx"
        if any(n.startswith("word/") for n in names):
            return "docx"
        if any(n.startswith("xl/") for n in names):
            return "xlsx"
        if any(n.startswith("ppt/") for n in names):
            return "pptx"
        if "content.xml" in names:
            return "odt"
        return "docx"

    # ── PDF ──────────────────────────────────────────────────────────

    def _extract_pdf(self, content: bytes) -> ExtractionResult:
        """Extract text from PDF with table preservation and OCR fallback."""
        if PDFPLUMBER_AVAILABLE:
            return self._extract_pdf_pdfplumber(content)
        if PYPDF_AVAILABLE:
            return self._extract_pdf_pypdf(content)
        logger.warning("No PDF library available (install pdfplumber or pypdf)")
        return ExtractionResult(error="PDF support is not installed on this server.")

    def _extract_pdf_pdfplumber(self, content: bytes) -> ExtractionResult:
        """Table-aware extraction using pdfplumber."""
        all_tables: list[str] = []
        page_texts: list[str] = []
        page_count = 0
        ocr = _OcrSession(content)

        try:
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                page_count = len(pdf.pages)

                for page in pdf.pages[:_MAX_PDF_PAGES]:
                    page_text = self._extract_pdfplumber_page(page, all_tables)

                    # OCR fallback for sparse pages
                    if len(page_text.strip()) < self.OCR_THRESHOLD:
                        page_text = ocr.improve(page.page_number - 1, page_text)

                    if page_text.strip():
                        page_texts.append(page_text)

        except Exception as e:  # pdfminer raises PDFSyntaxError and friends
            logger.error(f"pdfplumber extraction failed: {type(e).__name__}: {e}")
            ocr.close()
            # Fall back to pypdf
            if PYPDF_AVAILABLE:
                return self._extract_pdf_pypdf(content)
            raise ExtractionError(
                "The PDF could not be read. It may be corrupt or password-protected."
            ) from e
        finally:
            ocr.close()

        return self._pdf_result(page_texts, all_tables, page_count, ocr)

    @staticmethod
    def _pdf_result(
        page_texts: list[str], tables: list[str], page_count: int, ocr: "_OcrSession"
    ) -> ExtractionResult:
        """Assemble a PDF result, recording anything the caps left unread."""
        result = ExtractionResult(
            text="\n\n".join(page_texts),
            tables=tables,
            page_count=page_count,
            ocr_used=ocr.used,
        )
        if page_count > _MAX_PDF_PAGES:
            result.truncated = True
            result.warnings.append(
                f"Only the first {_MAX_PDF_PAGES:,} of {page_count:,} pages were read."
            )
        if ocr.skipped:
            result.truncated = True
            result.warnings.append(
                f"{ocr.skipped:,} scanned page(s) were not OCR'd (limit of "
                f"{_MAX_OCR_PAGES:,} OCR pages per document)."
            )
        return result

    def _extract_pdfplumber_page(self, page, all_tables: list[str]) -> str:
        """Extract text from a single pdfplumber page, preserving tables as markdown."""
        tables = page.find_tables()
        table_bboxes = [t.bbox for t in tables]

        # Extract text outside table regions
        if table_bboxes:
            # Build a list of non-table text by cropping outside table areas
            non_table_text = self._extract_text_outside_tables(page, table_bboxes)
        else:
            non_table_text = page.extract_text() or ""

        # Convert each table to markdown and collect
        md_tables: list[str] = []
        for table in tables:
            rows = table.extract()
            if rows:
                md = self._rows_to_markdown(rows)
                md_tables.append(md)
                all_tables.append(md)

        # Interleave: non-table text, then tables
        parts = []
        if non_table_text.strip():
            parts.append(non_table_text.strip())
        for md in md_tables:
            parts.append(md)

        return "\n\n".join(parts)

    def _extract_text_outside_tables(self, page, table_bboxes: list[tuple]) -> str:
        """Extract text from page regions that don't overlap with any table."""
        # Get all words on the page
        words = page.extract_words()
        if not words:
            return ""

        non_table_words = []
        for w in words:
            # Check if word center falls inside any table bbox
            cx = (w["x0"] + w["x1"]) / 2
            cy = (w["top"] + w["bottom"]) / 2
            in_table = False
            for bbox in table_bboxes:
                # bbox = (x0, top, x1, bottom)
                if bbox[0] <= cx <= bbox[2] and bbox[1] <= cy <= bbox[3]:
                    in_table = True
                    break
            if not in_table:
                non_table_words.append(w)

        if not non_table_words:
            return ""

        # Reconstruct text from words — group by line (top coordinate)
        lines: dict[float, list] = {}
        for w in non_table_words:
            # Round top to cluster words on the same line
            line_key = round(w["top"], 1)
            if line_key not in lines:
                lines[line_key] = []
            lines[line_key].append(w)

        text_lines = []
        for key in sorted(lines.keys()):
            line_words = sorted(lines[key], key=lambda w: w["x0"])
            text_lines.append(" ".join(w["text"] for w in line_words))

        return "\n".join(text_lines)

    def _rows_to_markdown(self, rows: list[list]) -> str:
        """Convert a list of row data into a markdown table string."""
        if not rows:
            return ""

        # Clean cells
        cleaned = []
        for row in rows:
            cleaned.append([(cell or "").replace("\n", " ").strip() for cell in row])

        # Determine column count
        col_count = max(len(r) for r in cleaned)
        # Pad short rows
        for row in cleaned:
            while len(row) < col_count:
                row.append("")

        # Build markdown
        lines = []
        # Header row
        lines.append("| " + " | ".join(cleaned[0]) + " |")
        # Separator
        lines.append("| " + " | ".join("---" for _ in range(col_count)) + " |")
        # Data rows
        for row in cleaned[1:]:
            lines.append("| " + " | ".join(row) + " |")

        return "\n".join(lines)

    def _extract_pdf_pypdf(self, content: bytes) -> ExtractionResult:
        """Basic PDF extraction via pypdf (no table awareness)."""
        page_texts = []
        page_count = 0
        ocr = _OcrSession(content)

        try:
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise ExtractionError(
                    "The PDF is password-protected. Remove the password and upload it again."
                )
            page_count = len(reader.pages)

            for i in range(min(page_count, _MAX_PDF_PAGES)):
                page_text = reader.pages[i].extract_text() or ""

                # OCR fallback for sparse pages
                if len(page_text.strip()) < self.OCR_THRESHOLD:
                    page_text = ocr.improve(i, page_text)

                if page_text.strip():
                    page_texts.append(page_text)

        except ExtractionError:
            raise
        except Exception as e:  # pypdf raises PdfReadError and friends
            logger.error(f"pypdf extraction failed: {type(e).__name__}: {e}")
            raise ExtractionError(
                "The PDF could not be read. It may be corrupt or password-protected."
            ) from e
        finally:
            ocr.close()

        return self._pdf_result(page_texts, [], page_count, ocr)

    # ── DOCX ─────────────────────────────────────────────────────────

    def _extract_docx(self, content: bytes) -> ExtractionResult:
        """Extract text from DOCX with tables converted to markdown."""
        if not DOCX_AVAILABLE:
            logger.warning("python-docx not available")
            return ExtractionResult(error="DOCX support is not installed on this server.")

        if content[:8] == _OLE_MAGIC:
            # A legacy .doc saved with a .docx name (or a password-protected
            # DOCX, which is also an OLE container) — the doc extractor tells
            # them apart.
            return self._extract_doc(content)
        self._check_zip_container(content)
        doc = DocxDocument(io.BytesIO(content))

        all_tables: list[str] = []
        parts: list[str] = []

        # Walk through the document body in order to preserve
        # the interleaving of paragraphs and tables.
        for element in doc.element.body:
            tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag

            if tag == "p":
                text = element.text or ""
                # Collect all runs
                runs = element.findall(
                    ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
                )
                if runs:
                    text = "".join((r.text or "") for r in runs)
                if text.strip():
                    parts.append(text.strip())

            elif tag == "tbl":
                # Find this table object in doc.tables
                md = self._docx_table_element_to_markdown(element)
                if md:
                    parts.append(md)
                    all_tables.append(md)

        return ExtractionResult(
            text="\n\n".join(parts),
            tables=all_tables,
            page_count=0,  # DOCX doesn't have a reliable page count
        )

    def _docx_table_element_to_markdown(self, tbl_element) -> str:
        """Convert a DOCX table XML element to a markdown table."""
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        rows_data: list[list[str]] = []

        for tr in tbl_element.findall(f"{ns}tr"):
            cells: list[str] = []
            for tc in tr.findall(f"{ns}tc"):
                # Collect all text runs in the cell
                cell_texts = []
                for p in tc.findall(f"{ns}p"):
                    runs = p.findall(f".//{ns}t")
                    p_text = "".join((r.text or "") for r in runs)
                    if p_text.strip():
                        cell_texts.append(p_text.strip())
                cells.append(" ".join(cell_texts))
            rows_data.append(cells)

        if not rows_data:
            return ""

        return self._rows_to_markdown(rows_data)

    # ── Legacy DOC ───────────────────────────────────────────────────

    def _extract_doc(self, content: bytes) -> ExtractionResult:
        """Extract text from a Word 97-2003 binary document (MS-DOC).

        Reads the piece table from the OLE streams directly (olefile), so no
        converter or external process is needed. Formatting and tables are
        flattened to text; that is enough for search and analysis.
        """
        if content[:8] != _OLE_MAGIC:
            if content[:4] == b"PK\x03\x04":
                # A .docx saved with a .doc name.
                return self._extract_docx(content)
            if content[:5] == b"{\\rtf":
                return self._extract_rtf(content)
            raise ExtractionError(
                "The file is not a Word document. Save it as .docx or PDF and upload it again."
            )
        if not MSG_AVAILABLE:
            return ExtractionResult(
                error="Legacy .doc support is not installed on this server. "
                "Save the file as .docx or PDF and upload it again."
            )
        try:
            ole = olefile.OleFileIO(io.BytesIO(content))
        except (OSError, ValueError) as e:
            raise ExtractionError("The file is not a valid Word document.") from e
        try:
            if ole.exists("EncryptedPackage"):
                raise ExtractionError(
                    "The document is password-protected. Remove the password and upload it again."
                )
            if not ole.exists("WordDocument"):
                raise ExtractionError(
                    "This legacy Office file is not a Word document. Save it in a current "
                    "format (.docx, .xlsx, .pptx or PDF) and upload it again."
                )
            word = ole.openstream("WordDocument").read()
            if len(word) < 0x01AA:
                raise ExtractionError("The Word document is truncated or corrupt.")
            flags = int.from_bytes(word[0x0A:0x0C], "little")
            if flags & 0x0100:  # fEncrypted
                raise ExtractionError(
                    "The document is password-protected. Remove the password and upload it again."
                )
            table_name = "1Table" if flags & 0x0200 else "0Table"  # fWhichTblStm
            if not ole.exists(table_name):
                raise ExtractionError("The Word document is truncated or corrupt.")
            table = ole.openstream(table_name).read()
        finally:
            ole.close()

        ccp_text = int.from_bytes(word[0x4C:0x50], "little")  # main-document characters
        fc_clx = int.from_bytes(word[0x01A2:0x01A6], "little")
        lcb_clx = int.from_bytes(word[0x01A6:0x01AA], "little")
        text = _read_doc_piece_table(word, table[fc_clx : fc_clx + lcb_clx], ccp_text)
        return ExtractionResult(text=_clean_doc_text(text))

    # ── Spreadsheets ─────────────────────────────────────────────────

    def _extract_xlsx(self, content: bytes) -> ExtractionResult:
        """Extract XLSX: every sheet becomes a markdown table under its name."""
        if not OPENPYXL_AVAILABLE:
            logger.warning("openpyxl not available")
            return ExtractionResult(error="XLSX support is not installed on this server.")
        self._check_zip_container(content)
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)

        parts: list[str] = []
        all_tables: list[str] = []
        for ws in wb.worksheets:
            rows = [
                ["" if cell is None else str(cell) for cell in row]
                for row in ws.iter_rows(values_only=True)
            ]
            rows = [r for r in rows if any(c.strip() for c in r)]
            if not rows:
                continue
            md = self._rows_to_markdown(rows)
            all_tables.append(md)
            parts.append(f"## {ws.title}\n\n{md}")
        wb.close()
        return ExtractionResult(text="\n\n".join(parts), tables=all_tables)

    def _extract_xls(self, content: bytes) -> ExtractionResult:
        """Extract legacy XLS via xlrd: every sheet becomes a markdown table."""
        if not XLRD_AVAILABLE:
            logger.warning("xlrd not available")
            return ExtractionResult(error="XLS support is not installed on this server.")
        wb = xlrd.open_workbook(file_contents=content)

        parts: list[str] = []
        all_tables: list[str] = []
        for ws in wb.sheets():
            rows = [
                [str(ws.cell_value(r, c)).strip() for c in range(ws.ncols)] for r in range(ws.nrows)
            ]
            rows = [r for r in rows if any(r)]
            if not rows:
                continue
            md = self._rows_to_markdown(rows)
            all_tables.append(md)
            parts.append(f"## {ws.name}\n\n{md}")
        return ExtractionResult(text="\n\n".join(parts), tables=all_tables)

    def _extract_delimited(self, content: bytes, delimiter: str) -> ExtractionResult:
        """Extract CSV/TSV as a markdown table."""
        text = self._extract_txt(content).text
        if not text.strip():
            return ExtractionResult()
        try:
            rows = [row for row in csv.reader(io.StringIO(text), delimiter=delimiter) if row]
        except csv.Error as e:
            logger.error(f"Delimited-file extraction failed: {e}")
            return ExtractionResult(text=text)
        if not rows:
            return ExtractionResult(text=text)
        md = self._rows_to_markdown(rows)
        return ExtractionResult(text=md, tables=[md])

    # ── Presentations ────────────────────────────────────────────────

    def _extract_pptx(self, content: bytes) -> ExtractionResult:
        """Extract PPTX: per-slide text, tables as markdown, speaker notes."""
        if not PPTX_AVAILABLE:
            logger.warning("python-pptx not available")
            return ExtractionResult(error="PPTX support is not installed on this server.")
        self._check_zip_container(content)
        prs = Presentation(io.BytesIO(content))

        parts: list[str] = []
        all_tables: list[str] = []
        slide_count = 0
        for i, slide in enumerate(prs.slides, start=1):
            slide_count = i
            slide_parts: list[str] = [f"## Slide {i}"]
            for shape in slide.shapes:
                if shape.has_table:
                    rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                    if rows:
                        md = self._rows_to_markdown(rows)
                        all_tables.append(md)
                        slide_parts.append(md)
                elif shape.has_text_frame and shape.text_frame.text.strip():
                    slide_parts.append(shape.text_frame.text.strip())
            notes = getattr(slide, "notes_slide", None) if slide.has_notes_slide else None
            if notes is not None and notes.notes_text_frame.text.strip():
                slide_parts.append(f"Notes: {notes.notes_text_frame.text.strip()}")
            if len(slide_parts) > 1:
                parts.append("\n\n".join(slide_parts))
        return ExtractionResult(text="\n\n".join(parts), tables=all_tables, page_count=slide_count)

    # ── ODT ──────────────────────────────────────────────────────────

    def _extract_odt(self, content: bytes) -> ExtractionResult:
        """Extract OpenDocument text via stdlib zip + XML (no extra library)."""
        import xml.etree.ElementTree as ET

        # Bounds content.xml too: its declared size is part of the total.
        self._check_zip_container(content)
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as zf:
                xml_bytes = zf.read("content.xml")
            # Uploads are untrusted, but stdlib etree is safe here: it never
            # resolves external entities, and expat >= 2.4.1 (python:3.11-slim
            # ships newer) caps entity expansion, so billion-laughs and
            # quadratic-blowup payloads are rejected rather than expanded.
            root = ET.fromstring(xml_bytes)  # noqa: S314  # nosec B314
        except (zipfile.BadZipFile, KeyError, ET.ParseError, OSError) as e:
            logger.error(f"ODT extraction failed: {e}")
            raise ExtractionError(
                "The file is not a valid OpenDocument text file (content.xml is missing "
                "or unreadable)."
            ) from e

        text_ns = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
        parts: list[str] = []
        for elem in root.iter():
            if elem.tag in (f"{text_ns}p", f"{text_ns}h"):
                t = "".join(elem.itertext()).strip()
                if t:
                    parts.append(t)
        return ExtractionResult(text="\n\n".join(parts))

    # ── HTML ─────────────────────────────────────────────────────────

    def _extract_html(self, content: bytes) -> ExtractionResult:
        """Strip HTML to readable text via the stdlib parser."""
        raw = self._extract_txt(content).text
        return ExtractionResult(text=_html_to_text(raw))

    # ── Email ────────────────────────────────────────────────────────

    def _extract_eml(self, content: bytes) -> ExtractionResult:
        """Extract RFC-822 email: headers, body (plain preferred), attachment names."""
        import email
        from email import policy

        try:
            msg = email.message_from_bytes(content, policy=policy.default)
        except (ValueError, OSError) as e:
            logger.error(f"EML extraction failed: {e}")
            return self._extract_txt(content)

        parts: list[str] = []
        for header in ("From", "To", "Cc", "Date", "Subject"):
            if msg.get(header):
                parts.append(f"{header}: {msg.get(header)}")

        body = ""
        attachments: list[str] = []
        try:
            plain = msg.get_body(preferencelist=("plain",))
            if plain is not None:
                body = plain.get_content()
            else:
                html_part = msg.get_body(preferencelist=("html",))
                if html_part is not None:
                    body = _html_to_text(html_part.get_content())
            for att in msg.iter_attachments():
                name = att.get_filename()
                if name:
                    attachments.append(name)
        except (ValueError, KeyError, OSError, LookupError) as e:
            logger.warning(f"EML body extraction degraded: {e}")

        if body.strip():
            parts.append(body.strip())
        if attachments:
            parts.append("Attachments: " + ", ".join(attachments))
        return ExtractionResult(text="\n\n".join(parts))

    def _extract_msg(self, content: bytes) -> ExtractionResult:
        """Extract Outlook .msg (MS-OXMSG OLE container): headers, body,
        attachment names — read straight from the property streams."""
        if not MSG_AVAILABLE:
            logger.warning("olefile not available")
            return ExtractionResult(error="Outlook .msg support is not installed on this server.")
        try:
            ole = olefile.OleFileIO(io.BytesIO(content))
        except (OSError, ValueError) as e:
            logger.error(f"MSG extraction failed: {e}")
            raise ExtractionError("The file is not a valid Outlook .msg message.") from e

        def read_str(prop: str, storage: str = "") -> str:
            # 001F = UTF-16LE, 001E = 8-bit; MS-OXMSG property stream naming.
            for suffix, enc in (("001F", "utf-16-le"), ("001E", "cp1252")):
                name = f"{storage}__substg1.0_{prop}{suffix}"
                if ole.exists(name):
                    try:
                        raw = ole.openstream(name).read()
                        return raw.decode(enc, errors="ignore").strip("\x00").strip()
                    except (OSError, ValueError):
                        return ""
            return ""

        def read_sent_time() -> str:
            # ClientSubmitTime (0039) / MessageDeliveryTime (0E06) live as
            # FILETIME values in the fixed-width __properties stream.
            if not ole.exists("__properties_version1.0"):
                return ""
            from datetime import datetime, timedelta

            try:
                data = ole.openstream("__properties_version1.0").read()
            except (OSError, ValueError):
                return ""
            for off in range(32, len(data) - 15, 16):
                tag = int.from_bytes(data[off : off + 4], "little")
                if tag in (0x00390040, 0x0E060040):
                    filetime = int.from_bytes(data[off + 8 : off + 16], "little")
                    if filetime:
                        epoch = datetime(1601, 1, 1, tzinfo=UTC)
                        return (epoch + timedelta(microseconds=filetime // 10)).isoformat()
            return ""

        try:
            parts: list[str] = []
            sender_name = read_str("0C1A")
            sender_addr = read_str("5D01") or read_str("0C1F")
            sender = (
                f"{sender_name} <{sender_addr}>"
                if sender_name and sender_addr
                else (sender_name or sender_addr)
            )
            for label, value in (
                ("From", sender),
                ("To", read_str("0E04")),
                ("Cc", read_str("0E03")),
                ("Date", read_sent_time()),
                ("Subject", read_str("0037")),
            ):
                if value:
                    parts.append(f"{label}: {value}")

            body = read_str("1000")
            if not body and ole.exists("__substg1.0_10130102"):
                try:
                    html_raw = ole.openstream("__substg1.0_10130102").read()
                    body = _html_to_text(html_raw.decode("utf-8", errors="ignore"))
                except (OSError, ValueError):
                    body = ""
            if body.strip():
                parts.append(body.strip())

            attachments: list[str] = []
            seen: set[str] = set()
            for entry in ole.listdir():
                top = entry[0]
                if top.startswith("__attach_version1.0_") and top not in seen:
                    seen.add(top)
                    name = read_str("3707", f"{top}/") or read_str("3704", f"{top}/")
                    if name:
                        attachments.append(name)
            if attachments:
                parts.append("Attachments: " + ", ".join(attachments))

            return ExtractionResult(text="\n\n".join(parts))
        finally:
            ole.close()

    # ── RTF ──────────────────────────────────────────────────────────

    def _extract_rtf(self, content: bytes) -> ExtractionResult:
        """Convert RTF to plain text (falls back to raw decode)."""
        raw = self._extract_txt(content).text
        if not RTF_AVAILABLE:
            return ExtractionResult(text=raw)
        try:
            return ExtractionResult(text=rtf_to_text(raw))
        except (ValueError, KeyError, IndexError) as e:
            logger.warning(f"RTF conversion failed, using raw text: {e}")
            return ExtractionResult(text=raw)

    # ── TXT ──────────────────────────────────────────────────────────

    def _extract_plain(self, content: bytes) -> ExtractionResult:
        """Plain-text path for txt/md/json and unrecognised formats.

        Unlike ``_extract_txt`` this refuses binary content: indexing the
        mojibake of an unsupported binary file helps no one.
        """
        if content[:2] not in (b"\xff\xfe", b"\xfe\xff") and b"\x00" in content[:8192]:
            raise ExtractionError(
                "This file type is not supported: the file is binary, not text. "
                "Supported formats are PDF, Word, Excel, PowerPoint, OpenDocument, "
                "email (.eml/.msg), RTF, HTML, CSV and plain text."
            )
        return self._extract_txt(content)

    def _extract_txt(self, content: bytes) -> ExtractionResult:
        """Decode text: UTF-8, then UTF-16 only when a BOM says so, then cp1252.

        UTF-16 without a BOM is not attempted — almost any even-length byte
        string "decodes" as UTF-16, producing garbage. cp1252 is tried before
        latin-1 because latin-1 never fails and would shadow it (curly quotes
        and dashes in Windows-authored files live in the cp1252-only range).
        """
        if content[:3] == b"\xef\xbb\xbf":
            return ExtractionResult(text=content[3:].decode("utf-8", errors="replace"))
        if content[:2] in (b"\xff\xfe", b"\xfe\xff"):
            return ExtractionResult(text=content.decode("utf-16", errors="replace"))
        for encoding in ("utf-8", "cp1252"):
            try:
                return ExtractionResult(text=content.decode(encoding))
            except (UnicodeDecodeError, ValueError):
                continue
        # cp1252 leaves five byte values undefined; latin-1 maps every byte.
        return ExtractionResult(text=content.decode("latin-1"))


_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _read_doc_piece_table(word: bytes, clx: bytes, ccp_text: int) -> str:
    """Reassemble MS-DOC main-document text from the Clx piece table.

    The Clx is a run of Prc entries (type 0x01) followed by one Pcdt (type
    0x02) holding a PlcPcd: n+1 character positions, then n 8-byte piece
    descriptors. Each descriptor points at a run in the WordDocument stream,
    stored either as cp1252 bytes ("compressed") or UTF-16LE.
    """
    pos = 0
    while pos < len(clx) and clx[pos] == 0x01:
        pos += 3 + int.from_bytes(clx[pos + 1 : pos + 3], "little")
    if pos >= len(clx) or clx[pos] != 0x02:
        raise ExtractionError("The Word document is truncated or corrupt.")
    lcb = int.from_bytes(clx[pos + 1 : pos + 5], "little")
    plc = clx[pos + 5 : pos + 5 + lcb]
    if len(plc) < 16 or (len(plc) - 4) % 12:
        raise ExtractionError("The Word document is truncated or corrupt.")
    n = (len(plc) - 4) // 12
    cps = [int.from_bytes(plc[i * 4 : i * 4 + 4], "little") for i in range(n + 1)]

    parts: list[str] = []
    for i in range(n):
        start = cps[i]
        end = min(cps[i + 1], ccp_text) if ccp_text else cps[i + 1]
        if end <= start:
            continue
        pcd = plc[(n + 1) * 4 + i * 8 : (n + 1) * 4 + (i + 1) * 8]
        fc = int.from_bytes(pcd[2:6], "little")
        count = end - start
        if fc & 0x40000000:  # fCompressed: one cp1252 byte per character
            offset = (fc & 0x3FFFFFFF) // 2
            parts.append(word[offset : offset + count].decode("cp1252", errors="replace"))
        else:
            offset = fc & 0x3FFFFFFF
            parts.append(word[offset : offset + count * 2].decode("utf-16-le", errors="replace"))
    return "".join(parts)


def _clean_doc_text(text: str) -> str:
    """Turn MS-DOC control characters into plain text."""
    out: list[str] = []
    # One entry per open field: True while inside its instruction (not text),
    # False once its separator has been seen (the displayed result follows).
    fields: list[bool] = []
    for ch in text:
        if ch == "\x13":  # field begin
            fields.append(True)
        elif ch == "\x14":  # field separator
            if fields:
                fields[-1] = False
        elif ch == "\x15":  # field end
            if fields:
                fields.pop()
        elif any(fields):
            continue
        elif ch in "\r\x0b\x0c":  # paragraph, line and page breaks
            out.append("\n")
        elif ch == "\x07":  # table cell mark; a second one in a row ends the row
            if out and out[-1] == "\t":
                out[-1] = "\n"
            else:
                out.append("\t")
        elif ch == "\x1e":  # non-breaking hyphen
            out.append("-")
        elif ch == "\t" or ch >= " ":
            out.append(ch)
    lines = [ln.strip() for ln in "".join(out).split("\n")]
    return "\n".join(ln for ln in lines if ln)


class _OcrSession:
    """OCR fallback for one PDF: opens the document once and enforces the page cap.

    Re-opening the PDF for every sparse page made a scanned document cost
    O(pages²); rendering is also bounded so one oversized page cannot exhaust
    memory.
    """

    def __init__(self, pdf_content: bytes):
        self._content = pdf_content
        self._doc = None
        self._pages_done = 0
        self.used = False
        self.skipped = 0

    def improve(self, page_index: int, page_text: str) -> str:
        """Return OCR text for a sparse page when it beats what the parser found."""
        if not OCR_AVAILABLE:
            return page_text
        if self._pages_done >= _MAX_OCR_PAGES:
            self.skipped += 1
            return page_text
        self._pages_done += 1
        ocr_text = self._ocr_page(page_index)
        if len(ocr_text.strip()) > len(page_text.strip()):
            self.used = True
            return ocr_text
        return page_text

    def _ocr_page(self, page_index: int) -> str:
        try:
            if self._doc is None:
                self._doc = fitz.open(stream=self._content, filetype="pdf")
            if page_index >= len(self._doc):
                return ""
            page = self._doc[page_index]
            scale = 300 / 72  # 300 DPI
            longest = max(page.rect.width, page.rect.height) * scale
            if longest > _MAX_OCR_PIXELS:
                scale *= _MAX_OCR_PIXELS / longest
            pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale))
            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            text = pytesseract.image_to_string(img)
            logger.debug(f"[OCR] Page {page_index + 1}: {len(text.strip())} chars")
            return text.strip()
        except Exception as e:  # PyMuPDF/tesseract errors must not fail the document
            logger.error(f"[OCR] Failed on page {page_index + 1}: {type(e).__name__}: {e}")
            return ""

    def close(self) -> None:
        if self._doc is not None:
            try:
                self._doc.close()
            except Exception:  # nosec B110 - best-effort cleanup
                pass
            self._doc = None


class _HTMLTextExtractor(HTMLParser):
    """Collects visible text; block-level tags become line breaks."""

    _SKIP = {"script", "style", "head", "title", "meta", "link"}
    _BLOCK = {
        "p",
        "div",
        "br",
        "li",
        "tr",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "section",
        "article",
        "header",
        "footer",
        "table",
        "blockquote",
    }

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self._chunks: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP:
            self._skip_depth += 1
        elif tag in self._BLOCK:
            self._chunks.append("\n")

    def handle_endtag(self, tag):
        if tag in self._SKIP and self._skip_depth > 0:
            self._skip_depth -= 1
        elif tag in self._BLOCK:
            self._chunks.append("\n")

    def handle_data(self, data):
        if self._skip_depth == 0 and data.strip():
            self._chunks.append(data)

    def text(self) -> str:
        raw = "".join(self._chunks)
        lines = [ln.strip() for ln in raw.splitlines()]
        return "\n".join(ln for ln in lines if ln)


def _html_to_text(html: str) -> str:
    parser = _HTMLTextExtractor()
    try:
        parser.feed(html)
        parser.close()
    except (ValueError, AssertionError) as e:
        logger.warning(f"HTML parse degraded: {e}")
    return parser.text()


# Module-level singleton
text_extraction_service = TextExtractionService()
