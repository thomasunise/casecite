"""
Shared "AI-generated — verify before relying" notice for exported artifacts.

Every generated document that leaves the app (analysis reports, redlines,
drafts, conversation transcripts) must end with this notice so a reader who
receives the file without the UI still knows it is machine output that has
not been verified.
"""

from __future__ import annotations

from typing import Any

AI_REVIEW_NOTICE = (
    "AI-generated draft. Verify every statement, citation and clause against the "
    "source documents and current law before relying on it or sending it to anyone."
)


def markdown_notice() -> list[str]:
    """Markdown lines for the notice: a rule, then the notice in italics."""
    return ["---", "", f"_{AI_REVIEW_NOTICE}_"]


def add_docx_notice(doc: Any) -> None:
    """Append a rule and an italic notice paragraph to a python-docx Document.

    Appends exactly two paragraphs at the end of the body; callers that read a
    fixed number of leading paragraphs (e.g. redline fidelity checks) are not
    affected.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rule = doc.add_paragraph()
    p_pr = rule._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "999999")
    borders.append(bottom)
    p_pr.append(borders)

    notice = doc.add_paragraph()
    notice.add_run(AI_REVIEW_NOTICE).italic = True


def pdf_notice_flowables(body_style: Any) -> list[Any]:
    """reportlab flowables for the notice: a rule, then the notice in italics."""
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import HRFlowable, Paragraph, Spacer

    style = ParagraphStyle("AINotice", parent=body_style, fontName="Helvetica-Oblique")
    return [
        Spacer(1, 10),
        HRFlowable(width="100%", thickness=0.5, color="#999999"),
        Spacer(1, 6),
        Paragraph(AI_REVIEW_NOTICE, style),
    ]
