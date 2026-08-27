"""
Conversation export — render a chat transcript as Markdown, Word or PDF.

The frontend sends the transcript it is displaying (every chat page keeps its
own message shape client-side, so the server does not reconstruct it from a
session). Assistant text is Markdown; a light block parser turns headings,
bullets and numbered items into document structure and strips inline markers
so Word/PDF readers don't see literal asterisks.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import UTC, datetime
from typing import Any

from app.services.ai_notice import add_docx_notice, markdown_notice, pdf_notice_flowables

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MIME = "application/pdf"
MD_MIME = "text/markdown; charset=utf-8"

FORMATS = {"docx": (DOCX_MIME, "docx"), "pdf": (PDF_MIME, "pdf"), "md": (MD_MIME, "md")}


class PdfUnavailable(RuntimeError):
    """reportlab is not installed in this deployment."""


@dataclass
class Citation:
    label: str
    quote: str | None = None
    url: str | None = None


@dataclass
class Message:
    role: str  # user | assistant
    text: str
    citations: list[Citation] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Markdown block parsing (assistant messages are Markdown)
# ---------------------------------------------------------------------------

_INLINE = [
    (re.compile(r"\*\*(.+?)\*\*"), r"\1"),
    (re.compile(r"__(.+?)__"), r"\1"),
    (re.compile(r"(?<!\w)\*(?!\s)(.+?)(?<!\s)\*(?!\w)"), r"\1"),
    (re.compile(r"(?<!\w)_(?!\s)(.+?)(?<!\s)_(?!\w)"), r"\1"),
    (re.compile(r"`([^`]+)`"), r"\1"),
    (re.compile(r"\[([^\]]+)\]\([^)]+\)"), r"\1"),
]
_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*•]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*(\d+)[.)]\s+(.*)$")


def strip_inline(text: str) -> str:
    for pattern, repl in _INLINE:
        text = pattern.sub(repl, text)
    return text


def parse_blocks(text: str) -> list[tuple[str, str]]:
    """Return (kind, text) blocks: heading / bullet / number / para."""
    blocks: list[tuple[str, str]] = []
    for raw in (text or "").replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        if m := _HEADING.match(line):
            blocks.append(("heading", strip_inline(m.group(2).strip())))
        elif m := _BULLET.match(line):
            blocks.append(("bullet", strip_inline(m.group(1).strip())))
        elif m := _NUMBERED.match(line):
            blocks.append(("number", strip_inline(m.group(2).strip())))
        else:
            blocks.append(("para", strip_inline(line.strip())))
    return blocks


def _coerce(messages: list[Any]) -> list[Message]:
    out: list[Message] = []
    for m in messages:
        get = m.get if isinstance(m, dict) else lambda k, _m=m: getattr(_m, k, None)
        cites = []
        for c in get("citations") or []:
            cget = c.get if isinstance(c, dict) else lambda k, _c=c: getattr(_c, k, None)
            label = str(cget("label") or "").strip()
            if not label:
                continue
            cites.append(
                Citation(label=label, quote=cget("quote") or None, url=cget("url") or None)
            )
        out.append(
            Message(
                role="user" if str(get("role") or "").lower() == "user" else "assistant",
                text=str(get("text") or ""),
                citations=cites,
            )
        )
    return out


def _stamp() -> str:
    return datetime.now(UTC).strftime("%d %B %Y, %H:%M UTC")


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------


def render_markdown(title: str, messages: list[Any]) -> str:
    lines = [f"# {title or 'Conversation'}", "", f"_Exported {_stamp()}_", ""]
    for msg in _coerce(messages):
        lines.append("## You" if msg.role == "user" else "## Assistant")
        lines.append("")
        lines.append(msg.text.strip())
        lines.append("")
        if msg.citations:
            lines.append("**Sources**")
            lines.append("")
            for i, c in enumerate(msg.citations, 1):
                head = f"{i}. {c.label}"
                if c.url:
                    head += f" — {c.url}"
                lines.append(head)
                if c.quote:
                    lines.append(f"   > {c.quote.strip()}")
            lines.append("")
    lines += markdown_notice()
    return "\n".join(lines).rstrip() + "\n"


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------


def render_docx(title: str, messages: list[Any]) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(title or "Conversation", level=1)
    stamp = doc.add_paragraph(f"Exported {_stamp()}")
    stamp.runs[0].italic = True

    for msg in _coerce(messages):
        doc.add_heading("You" if msg.role == "user" else "Assistant", level=2)
        if msg.role == "user":
            for raw in msg.text.replace("\r\n", "\n").split("\n"):
                if raw.strip():
                    doc.add_paragraph(strip_inline(raw.strip()))
        else:
            for kind, text in parse_blocks(msg.text):
                if kind == "heading":
                    doc.add_heading(text, level=3)
                elif kind == "bullet":
                    doc.add_paragraph(text, style="List Bullet")
                elif kind == "number":
                    doc.add_paragraph(text, style="List Number")
                else:
                    doc.add_paragraph(text)
        if msg.citations:
            p = doc.add_paragraph()
            p.add_run("Sources").bold = True
            for i, c in enumerate(msg.citations, 1):
                line = doc.add_paragraph(style="List Number")
                line.add_run(c.label)
                if c.url:
                    line.add_run(f" — {c.url}")
                if c.quote:
                    q = doc.add_paragraph()
                    q.paragraph_format.left_indent = Pt(24)
                    run = q.add_run(c.quote.strip())
                    run.italic = True
                    run.font.size = Pt(10)

    add_docx_notice(doc)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def _esc(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def render_pdf(title: str, messages: list[Any]) -> bytes:
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer
    except ImportError as e:  # pragma: no cover - depends on deployment
        raise PdfUnavailable("PDF export requires reportlab.") from e

    styles = getSampleStyleSheet()
    body = ParagraphStyle("Body", parent=styles["Normal"], fontSize=11, leading=15, spaceAfter=5)
    quote = ParagraphStyle(
        "Quote", parent=body, fontSize=9.5, leading=13, leftIndent=18, textColor="#444444"
    )
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=18, leading=22, spaceAfter=4)
    h2 = ParagraphStyle(
        "H2", parent=styles["Heading2"], fontSize=12.5, leading=16, spaceBefore=12, spaceAfter=4
    )
    h3 = ParagraphStyle(
        "H3", parent=styles["Heading3"], fontSize=11.5, leading=15, spaceBefore=8, spaceAfter=3
    )
    meta = ParagraphStyle("Meta", parent=body, fontSize=9, textColor="#666666", spaceAfter=10)

    flow: list[Any] = [
        Paragraph(_esc(title or "Conversation"), h1),
        Paragraph(f"Exported {_stamp()}", meta),
    ]

    for msg in _coerce(messages):
        flow.append(Paragraph("You" if msg.role == "user" else "Assistant", h2))
        if msg.role == "user":
            for raw in msg.text.replace("\r\n", "\n").split("\n"):
                if raw.strip():
                    flow.append(Paragraph(_esc(strip_inline(raw.strip())), body))
        else:
            items: list[Any] = []
            numbered: bool | None = None

            def flush() -> None:
                nonlocal items, numbered
                if items:
                    flow.append(
                        ListFlowable(
                            items,
                            bulletType="1" if numbered else "bullet",
                            leftIndent=16,
                            bulletFontSize=9,
                        )
                    )
                items, numbered = [], None

            for kind, text in parse_blocks(msg.text):
                if kind in ("bullet", "number"):
                    is_num = kind == "number"
                    if numbered is not None and numbered != is_num:
                        flush()
                    numbered = is_num
                    items.append(ListItem(Paragraph(_esc(text), body), leftIndent=16))
                    continue
                flush()
                flow.append(Paragraph(_esc(text), h3 if kind == "heading" else body))
            flush()
        if msg.citations:
            flow.append(Paragraph("<b>Sources</b>", body))
            for i, c in enumerate(msg.citations, 1):
                head = f"{i}. {_esc(c.label)}"
                if c.url:
                    head += f" — {_esc(c.url)}"
                flow.append(Paragraph(head, body))
                if c.quote:
                    flow.append(Paragraph(_esc(c.quote.strip()), quote))
        flow.append(Spacer(1, 6))

    flow.extend(pdf_notice_flowables(body))

    buf = io.BytesIO()
    SimpleDocTemplate(
        buf,
        pagesize=letter,
        leftMargin=1 * inch,
        rightMargin=1 * inch,
        topMargin=1 * inch,
        bottomMargin=1 * inch,
        title=title or "Conversation",
    ).build(flow)
    return buf.getvalue()


def render(fmt: str, title: str, messages: list[Any]) -> bytes:
    if fmt == "docx":
        return render_docx(title, messages)
    if fmt == "pdf":
        return render_pdf(title, messages)
    if fmt == "md":
        return render_markdown(title, messages).encode("utf-8")
    raise ValueError(f"Unsupported export format: {fmt}")


def safe_filename(title: str, ext: str) -> str:
    ascii_title = (title or "").encode("ascii", "ignore").decode()
    stem = "".join(c for c in ascii_title if c.isalnum() or c in " -_").strip()
    stem = re.sub(r"\s+", "-", stem)[:60].strip("-") or "conversation"
    return f"{stem}.{ext}"
