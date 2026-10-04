"""
Legal-document support endpoints still used by the live UI.

Only the endpoints the source-available build actually calls are kept here:
- GET  /legal-docs/courts        — court list for jurisdiction pickers (Contracts)
- POST /legal-docs/extract-text  — text extraction from an uploaded file
- POST /legal-docs/convert-to-pdf — render an uploaded doc to PDF for the viewer

(The former drafting/templates/discovery routers were removed; these three are
the only pieces live features — Contracts and Case Citations — depend on.)
"""

import asyncio
import io
import logging

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import StreamingResponse

from app.config import settings
from app.services.auth import TokenData, get_current_user
from app.services.courts import get_all_courts
from app.services.permissions import require_any_permission
from app.utils.upload_validation import read_upload_capped, validate_import_file

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/legal-docs", tags=["legal-documents"])

# Extraction and conversion feed the document viewer, Contracts and Case
# Citations, so any one of those capabilities is enough to call them.
_FILE_TOOL_PERMISSIONS = ("documents.view", "contracts.use", "authority_map.use")


@router.get("/courts")
async def get_courts(
    current_user: TokenData = Depends(get_current_user),
) -> dict:
    """Get the static list of courts used by jurisdiction pickers.

    Requires authentication (same policy as the /tools research endpoints).
    """
    courts = get_all_courts()
    return {"courts": courts, "count": len(courts)}


async def _extract(content: bytes, filename: str):
    """Extract text off the event loop; a file with no usable text is a 422
    carrying the reason (scanned PDF without OCR, corrupt archive, ...)."""
    from app.services.text_extraction import text_extraction_service

    result = await text_extraction_service.extract_async(file_content=content, filename=filename)
    if result.error and not result.text.strip():
        raise HTTPException(status_code=422, detail=result.error)
    return result


@router.post("/extract-text")
async def extract_text_from_file(
    request: Request,
    file: UploadFile = File(...),
    current_user: TokenData = require_any_permission(*_FILE_TOOL_PERMISSIONS),
) -> dict:
    """Extract text content from an uploaded file (PDF, DOCX, DOC, TXT, RTF, MD).

    Requires a document, contract or authority-map permission. The upload is run
    through the shared allowlist + per-type size cap + magic-byte check before
    it reaches the (expensive, OCR-capable) extraction pipeline, which runs in a
    worker thread so one large file cannot stall the event loop. A file that
    yields no text returns 422 with the reason.
    """
    try:
        # Chunked read with a hard cap — an oversized body is rejected before
        # it is buffered into memory (413), then validated per-type below.
        content = await read_upload_capped(file, settings.max_upload_size, request=request)
        # Reject oversized / disallowed / spoofed files before any parsing.
        safe_name, _ = validate_import_file(content, file.content_type, file.filename or "")
        result = await _extract(content, safe_name)
        return {
            "success": True,
            "filename": file.filename,
            "text": result.text,
            "char_count": result.char_count,
            "word_count": result.word_count,
            "page_count": result.page_count,
            "tables_found": len(result.tables),
            "ocr_used": result.ocr_used,
            "truncated": result.truncated,
            "warnings": result.warnings,
        }
    except HTTPException:
        raise
    except (ValueError, OSError, UnicodeDecodeError) as e:
        logger.error(f"Failed to extract text: {e}", exc_info=True)
        raise HTTPException(
            status_code=400,
            detail="Failed to extract text. Please upload a valid PDF, DOCX, TXT, or Markdown file.",
        )


def _escape(text: str) -> str:
    """Escape text for a reportlab Paragraph (which parses inline markup)."""
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _render_pdf(blocks: list[tuple[str, bool]]) -> bytes:
    """Render (text, is_heading) blocks to a letter-size PDF; "" is a blank line."""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    body_style = ParagraphStyle(
        "DocBody", parent=styles["Normal"], fontSize=12, leading=16, spaceAfter=6
    )
    heading_style = ParagraphStyle(
        "DocHeading",
        parent=styles["Heading2"],
        fontSize=14,
        leading=18,
        spaceBefore=12,
        spaceAfter=6,
    )
    flowables = []
    for text, is_heading in blocks:
        if not text.strip():
            flowables.append(Spacer(1, 6))
        else:
            flowables.append(Paragraph(_escape(text), heading_style if is_heading else body_style))
    if not flowables:
        flowables.append(Paragraph("(empty document)", body_style))

    pdf_buffer = io.BytesIO()
    SimpleDocTemplate(
        pdf_buffer,
        pagesize=letter,
        leftMargin=1 * inch,
        rightMargin=1 * inch,
        topMargin=1 * inch,
        bottomMargin=1 * inch,
    ).build(flowables)
    return pdf_buffer.getvalue()


def _docx_to_pdf(content: bytes) -> bytes:
    """Blocking DOCX -> PDF (python-docx + reportlab) — call via asyncio.to_thread."""
    from docx import Document as DocxDocument

    blocks = []
    for para in DocxDocument(io.BytesIO(content)).paragraphs:
        style_name = (para.style.name or "").lower()
        blocks.append((para.text.strip(), "heading" in style_name or "title" in style_name))
    return _render_pdf(blocks)


def _pdf_response(pdf: bytes) -> StreamingResponse:
    return StreamingResponse(
        io.BytesIO(pdf),
        media_type="application/pdf",
        headers={"Content-Disposition": "inline; filename=document.pdf"},
    )


@router.post("/convert-to-pdf")
async def convert_to_pdf(
    request: Request,
    file: UploadFile = File(...),
    current_user: TokenData = require_any_permission(*_FILE_TOOL_PERMISSIONS),
):
    """Convert an uploaded document to PDF for the visual viewer.

    Requires a document, contract or authority-map permission. PDFs pass
    through; DOCX and text/markdown are rendered with reportlab in a worker
    thread. The upload is validated (allowlist + per-type size cap + magic
    bytes) before any parsing/rendering.
    """
    try:
        # Chunked read with a hard cap — an oversized body is rejected before
        # it is buffered into memory (413), then validated per-type below.
        content = await read_upload_capped(file, settings.max_upload_size, request=request)
        # Reject oversized / disallowed / spoofed files before any parsing.
        validate_import_file(content, file.content_type, file.filename or "")
        filename = (file.filename or "").lower()

        # PDF — pass through unchanged.
        if filename.endswith(".pdf") or content[:4] == b"%PDF":
            return _pdf_response(content)

        # DOCX — convert via python-docx + reportlab.
        if filename.endswith(".docx") or content[:4] == b"PK\x03\x04":
            return _pdf_response(await asyncio.to_thread(_docx_to_pdf, content))

        # TXT / RTF / Markdown / other text — render as a simple PDF.
        text = (await _extract(content, filename)).text
        blocks = [(line, False) for line in text.split("\n")]
        return _pdf_response(await asyncio.to_thread(_render_pdf, blocks))

    except HTTPException:
        raise
    except (ValueError, OSError, UnicodeDecodeError) as e:
        logger.error(f"Failed to convert to PDF: {e}", exc_info=True)
        raise HTTPException(status_code=400, detail="Failed to convert file to PDF.")
